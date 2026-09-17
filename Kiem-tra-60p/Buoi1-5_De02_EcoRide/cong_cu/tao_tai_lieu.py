"""Dựng đề thi và đáp án (PDF + DOCX) cho Đề số 02 — EcoRide.

Chạy (sau tao_du_lieu.py và sau khi đã chạy notebook đáp án để có dap_an/bieu_do.png):
    python3 tao_tai_lieu.py

Nội dung viết một lần dưới dạng danh sách khối, rồi xuất ra:
  - HTML → PDF bằng Google Chrome headless
  - DOCX bằng python-docx (bản để giảng viên sửa tên trường, học kỳ…)
Các con số trong đáp án ứng với SEED = 2502 trong tao_du_lieu.py.
"""
import html
import re
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GOC = Path(__file__).resolve().parent.parent
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'
DE_SO = '02'
TEN_MON = 'Lập trình phân tích dữ liệu'


# =====================================================================================
# NỘI DUNG — cú pháp nội dòng: **đậm**, `mã`, ~nghiêng~
# =====================================================================================

def h2(ten, diem=''):
    return ('h2', ten, diem)


def h3(ten, diem=''):
    return ('h3', ten, diem)


def p(text):
    return ('p', text)


def ul(*items):
    return ('ul', items)


def ol(*items):
    return ('ol', items)


def code(text):
    return ('code', text.strip('\n'))


def table(header, rows, widths=None, small=False):
    return ('table', header, rows, widths, small)


def note(nhan, text):
    return ('note', nhan, text)


def ex(text):
    """Dòng ví dụ định dạng (khối mã một dòng)."""
    return ('code', text)


def de_thi():
    return [
        ('header', 'BÀI KIỂM TRA THỰC HÀNH', None),
        ('title', TEN_MON.upper(), f'Đề số {DE_SO} · Phạm vi Buổi 1–5 · Bộ dữ liệu: Dịch vụ xe công cộng EcoRide'),
        ('info', [('THỜI GIAN LÀM BÀI', '60 phút (không kể thời gian phát đề)'),
                  ('TỔNG ĐIỂM', '10 điểm — 5 phần, 11 câu'),
                  ('HÌNH THỨC', 'Thực hành trên máy · Jupyter Notebook hoặc Google Colab'),
                  ('THƯ VIỆN ĐƯỢC DÙNG', 'Python chuẩn (re, collections, sqlite3), numpy, pandas, matplotlib, '
                                         'seaborn, scipy, scikit-learn'),
                  ('TÀI LIỆU', 'Được dùng tài liệu cá nhân'),
                  ('KHÔNG ĐƯỢC', 'Dùng Internet, công cụ AI, trao đổi bài')]),
        ('sv',),
        h2('Hướng dẫn làm bài & nộp bài'),
        ol('Làm bài trong notebook `bai_lam/bai_lam_mau.ipynb`, đổi tên thành `MSSV_HoTen.ipynb`. '
           'Mỗi câu làm trong đúng ô được đánh dấu.',
           'Kết quả phải được in ra trong notebook (bằng `print` hoặc hiển thị bảng). Câu hỏi yêu cầu trả lời '
           'bằng lời thì viết vào ô Markdown “Trả lời”.',
           'Trước khi nộp: chọn **Restart & Run All**. Ô báo lỗi hoặc chưa chạy sẽ không được chấm.',
           'Nộp 3 tập tin: notebook bài làm, `dong_loi.txt` (câu A1) và `bieu_do.png` (Phần E).'),
        p('Phân bổ thời gian gợi ý: Phần A 10 phút · Phần B 10 phút · Phần C 18 phút · Phần D 10 phút · '
          'Phần E 12 phút.'),
        table(['Buổi', 'Nội dung được kiểm tra', 'Câu'], [
            ['1', 'Python cơ bản · khối lệnh & cấu trúc điều khiển · hàm tự định nghĩa', 'A1, A2'],
            ['2', 'list, tuple, dict, set · mảng & ma trận NumPy · đại số ma trận', 'A2, B1, B2'],
            ['3', 'Tập tin văn bản · tập tin dữ liệu với Pandas · truy cập CSDL · lập trình hướng đối tượng',
             'A1, B3, C1'],
            ['4', 'Tiền xử lý: làm sạch dữ liệu · tích hợp dữ liệu · EDA', 'C2, C3, E'],
            ['5', 'Tiền xử lý: biến đổi dữ liệu · thu gọn dữ liệu', 'D1, D2'],
        ], widths=[7, 80, 13], small=True),

        h2('Bối cảnh & dữ liệu'),
        p('**EcoRide** cho thuê xe đạp, xe đạp điện và xe máy điện qua ứng dụng tại 6 trạm ở Đà Nẵng và Hội An. '
          'Khách mở khoá xe ở một trạm, đi rồi trả xe ở một trạm bất kỳ trong cùng thành phố; cước tính theo phút. '
          'Bạn nhận dữ liệu từ tháng 3 đến tháng 8/2025 trong thư mục `du_lieu/` gồm 4 tập tin. '
          'Dữ liệu là dữ liệu thô, có lỗi — đó là một phần của bài.'),
        table(['Tập tin', 'Nội dung'], [
            ['`nhat_ky_khoa_xe.txt`', 'Nhật ký khoá thông minh của xe điện tháng 7/2025. Mỗi lần mở khoá (MUON) '
                                      'hoặc khoá lại (TRA) ghi một dòng kèm mức pin lúc đó.'],
            ['`chuyen_di.csv`', 'Mỗi dòng là một chuyến đi. Các cột mô tả ở bảng dưới.'],
            ['`khach_hang.xlsx`', 'Sheet `khach_hang`: ma_kh, nam_sinh, so_dien_thoai, goi_cuoc, ngay_dang_ky<br>'
                                  'Sheet `goi_cuoc`: goi_cuoc, giam_gia (tỉ lệ, vd. 0.20 = giảm 20%)'],
            ['`he_thong.db`', 'Cơ sở dữ liệu SQLite, 2 bảng. `tram`: id, ten_tram, thanh_pho, suc_chua (số chỗ đỗ). '
                              '`loai_xe`: ma_loai, ten_loai, phi_mo_khoa, gia_phut, gia_30_phut (VNĐ)'],
        ], widths=[24, 76]),
        table(['Cột trong chuyen_di.csv', 'Ý nghĩa'], [
            ['ma_cd', 'Mã chuyến đi, dạng CD + 5 chữ số'],
            ['bat_dau', 'Thời điểm mở khoá xe, định dạng ngày/tháng/năm giờ:phút (dd/mm/yyyy HH:MM)'],
            ['ket_thuc', 'Thời điểm trả xe, cùng định dạng; có thể trống'],
            ['tram_muon · tram_tra', 'Mã trạm mượn xe · mã trạm trả xe'],
            ['ma_kh', 'Mã khách có đăng ký gói cước; để trống nếu là khách lẻ (quét mã, trả tiền từng chuyến)'],
            ['loai_xe', 'Xe đạp, Xe đạp điện, Xe máy điện'],
            ['quang_duong_km', 'Quãng đường GPS ghi nhận được (km)'],
            ['danh_gia', 'Rất không hài lòng, Không hài lòng, Bình thường, Hài lòng, Rất hài lòng; có thể trống'],
        ], widths=[24, 76]),

        h2('Phần A — Python thuần: tập tin văn bản, regex, hàm', '2,0 điểm'),
        p('Phần này **không dùng pandas và numpy**. Chỉ dùng Python chuẩn: `open`, `re`, `collections`…'),
        p('Một dòng hợp lệ trong `nhat_ky_khoa_xe.txt` có đúng dạng:'),
        ex('2025-07-02 07:15 | DN03 | XD014 | TRA | pin=18%'),
        ul('Ngày dạng YYYY-MM-DD và giờ dạng HH:MM, cách nhau đúng một khoảng trắng.',
           'Mã trạm là DN (Đà Nẵng) hoặc HA (Hội An) + đúng 2 chữ số; mã xe là XD (xe đạp điện) hoặc '
           'XM (xe máy điện) + đúng 3 chữ số.',
           'Sự kiện là một trong hai giá trị viết hoa: MUON (mở khoá), TRA (khoá lại).',
           'Trường cuối dạng `pin=N%` với N gồm 1–3 chữ số; dòng **kết thúc ngay sau** dấu %.',
           'Các trường ngăn cách bởi ` | ` (khoảng trắng – gạch đứng – khoảng trắng).',
           'Dòng trống và dòng bắt đầu bằng # được bỏ qua — không tính là hợp lệ, cũng không tính là lỗi.'),
        h3('Câu A1. Lọc dòng hợp lệ bằng biểu thức chính quy', '1,0 điểm'),
        p('Đọc tập tin, dùng module `re` phân loại từng dòng thành hợp lệ hoặc không hợp lệ.'),
        ol('In số dòng hợp lệ và số dòng không hợp lệ.',
           'Dùng `Counter`, in số sự kiện của từng trạm trong các dòng hợp lệ, sắp xếp theo mã trạm.',
           'Ghi toàn bộ dòng không hợp lệ (giữ nguyên nội dung, mỗi dòng một dòng) vào tập tin `dong_loi.txt`.'),
        h3('Câu A2. Thống kê pin yếu', '1,0 điểm'),
        p('Viết hàm `ti_le_pin_yeu(cac_dong, nguong=20, toi_thieu=8)` nhận danh sách dòng hợp lệ và trả về dict '
          '`{mã xe: tỉ lệ pin yếu}`. **Chỉ xét sự kiện TRA.** Một lần trả xe bị tính là **pin yếu** khi mức pin '
          '**nhỏ hơn hẳn** `nguong` (trả xe với pin=20% không tính là pin yếu). Tỉ lệ = số lần trả pin yếu ÷ '
          'tổng số lần trả của xe đó. Chỉ đưa vào kết quả những xe có ít nhất `toi_thieu` lần trả.'),
        ul('Gọi hàm với tham số mặc định, in tỉ lệ của từng xe (3 chữ số thập phân) và cho biết xe có tỉ lệ '
           'pin yếu cao nhất.',
           'Dùng kiểu `set`, liệt kê các xe đã có sự kiện ở **cả hai thành phố** (xét tất cả dòng hợp lệ; '
           'thành phố là 2 ký tự đầu của mã trạm).'),

        h2('Phần B — NumPy & lập trình hướng đối tượng', '2,0 điểm'),
        h3('Câu B1. Ma trận lượt thuê', '0,5 điểm'),
        p('Ma trận `L` (đã có sẵn trong notebook) là số lượt thuê xe trong một ngày thứ Bảy tại 6 trạm '
          'DN01, DN02, DN03, DN04, HA01, HA02 (dòng) theo 5 khung giờ 05–08h, 08–11h, 11–14h, 14–17h, 17–20h '
          '(cột). **Không dùng vòng lặp.**'),
        code('''
L = np.array([[42, 35, 28, 39, 71],
              [58, 31, 22, 34, 66],
              [25, 47, 44, 41, 52],
              [18, 26, 19, 23, 37],
              [12, 38, 41, 45, 83],
              [33, 29, 24, 36, 48]])
'''),
        ol('Tổng lượt thuê của từng trạm và trạm cao nhất; trung bình theo từng khung giờ và khung giờ có '
           'trung bình cao nhất.',
           'Tỉ trọng lượt thuê **giờ cao điểm** (khung 05–08h và khung 17–20h) trong tổng lượt của từng trạm — '
           'dùng broadcasting; dùng mặt nạ boolean in ra các trạm có tỉ trọng này lớn hơn 0,5.',
           'Số ô có lượt thuê lớn hơn lượt thuê trung bình **của chính khung giờ đó** (trung bình trên 6 trạm).'),
        h3('Câu B2. Giải hệ phương trình', '0,5 điểm'),
        p('Cước một chuyến xe máy điện được tính theo công thức: cước = phí mở khoá + giá mỗi phút × số phút + '
          'phụ phí mỗi km × số km. Ba tham số giá chưa biết; cước của ba chuyến như sau:'),
        table(['Chuyến', 'Số phút', 'Số km', 'Cước (nghìn đồng)'],
              [['1', '20', '3', '23'], ['2', '35', '8', '42'], ['3', '50', '9', '53']], widths=[16, 20, 20, 44]),
        p('Lập ma trận hệ số và vector vế phải, dùng `np.linalg.solve` tìm ba tham số giá, kiểm tra lại nghiệm '
          'bằng phép nhân ma trận, rồi dùng nghiệm dự đoán cước của một chuyến 30 phút, 6 km.'),
        h3('Câu B3. Lớp chuyến đi', '1,0 điểm'),
        ul('Lớp `ChuyenDi(ma, phut, km, gia_phut)` có **thuộc tính lớp** `PHI_MO_KHOA = 5000`. `phut` là '
           '**property**; setter ném `ValueError` nếu thời lượng ≤ 0 (áp dụng cả lúc khởi tạo).',
           'Phương thức `cuoc_phi()` = PHI_MO_KHOA + phut × gia_phut; `toc_do_tb()` = km ÷ (phut ÷ 60), đơn vị '
           'km/h; `__str__` trả về chuỗi đúng mẫu `CD001: 25 phút · 6.0 km · 14.4 km/h · cước 17,500đ`.',
           'Lớp con `ChuyenDiGoiThang(ma, phut, km, gia_phut, phut_mien_phi=30)` kế thừa `ChuyenDi`, gọi '
           '`super().__init__`, ghi đè `cuoc_phi()`: không thu phí mở khoá, chỉ tính tiền các phút vượt quá '
           '`phut_mien_phi` (không bao giờ âm).'),
        p('Chạy kiểm thử: in `ChuyenDi(\'CD001\', 25, 6.0, 500)`; in `ChuyenDiGoiThang(\'CD002\', 48, 10.4, 500)`; '
          'gán `phut` của chuyến CD001 thành 0 trong khối `try/except` và in thông báo lỗi.'),

        h2('Phần C — Pandas: đọc, làm sạch và tích hợp dữ liệu', '2,5 điểm'),
        h3('Câu C1. Đọc dữ liệu từ ba loại nguồn', '0,5 điểm'),
        ul('Đọc `chuyen_di.csv`; chuyển **cả hai** cột `bat_dau`, `ket_thuc` sang kiểu datetime, chú ý định dạng '
           'ngày/tháng/năm.',
           'Đọc cả hai sheet của `khach_hang.xlsx`.',
           'Kết nối `he_thong.db` bằng `sqlite3`, đọc hai bảng bằng `pd.read_sql` với câu lệnh SQL.',
           'In kích thước (shape) của 5 bảng; in số giá trị thiếu của các cột có giá trị thiếu trong bảng '
           'chuyến đi.'),
        h3('Câu C2. Làm sạch bảng chuyến đi', '1,0 điểm'),
        p('Thực hiện **đúng thứ tự** năm bước sau:'),
        ol('Loại bỏ các dòng trùng lặp hoàn toàn.',
           'Chuẩn hoá `loai_xe` về đúng 3 giá trị `Xe đạp`, `Xe đạp điện`, `Xe máy điện`. Lưu ý dữ liệu có cả '
           'cách viết không dấu, và “xe đạp” là phần đầu của “xe đạp điện”.',
           'Tạo cột `phut` = số phút từ `bat_dau` đến `ket_thuc`. Đổi các giá trị `phut` ≤ 0 (lỗi đồng hồ) thành '
           'giá trị thiếu. Loại bỏ các chuyến có `phut` nhỏ hơn 2 (mở khoá thử) — các chuyến đang trống `phut` '
           'phải được giữ lại.',
           'Đổi các giá trị `quang_duong_km` âm hoặc lớn hơn 80 (lỗi GPS) thành giá trị thiếu.',
           'Tìm ngoại lai `phut` theo quy tắc IQR, với Q1, Q3 **tính riêng cho từng loại xe** (ngoại lai khi '
           '> Q3 + 1,5 × IQR). In số ngoại lai của mỗi loại xe, kéo các giá trị đó về đúng ngưỡng trên của loại '
           'xe (winsorization). Cuối cùng điền `phut` và `quang_duong_km` còn trống bằng **trung vị theo loại xe**.'),
        p('In ra: số dòng sau bước 1 và sau bước 3; số giá trị bị đổi thành thiếu ở bước 3 và ở bước 4; số ngoại '
          'lai theo loại xe; số chuyến theo từng loại xe sau khi làm sạch.'),
        h3('Câu C3. Tích hợp dữ liệu', '1,0 điểm'),
        ol('Bảng `loai_xe`: viết code chứng minh `gia_30_phut` là cột dẫn xuất (= phi_mo_khoa + 30 × gia_phut ở '
           'mọi dòng) rồi xoá cột này.',
           'Bảng khách hàng: in số mã khách bị trùng; loại trùng sao cho **mỗi `ma_kh` chỉ còn một dòng**. '
           'Chuẩn hoá `so_dien_thoai`: chỉ giữ chữ số, đầu số 84 đổi thành 0 (vd. `+84 905 123 456` → '
           '`0905123456`); in số khách có số điện thoại không đủ 10 chữ số. Ghép sheet `goi_cuoc` để mỗi khách '
           'có `giam_gia`.',
           'Loại các chuyến có `tram_muon` không có trong bảng `tram` (in số dòng bị loại). Ghép chuyến đi với '
           '`loai_xe` (khoá `loai_xe` ↔ `ten_loai`), khách hàng và `tram` (khoá `tram_muon` ↔ `id`) sao cho '
           '**không mất chuyến nào** còn lại; chuyến không có khách khớp: `goi_cuoc = "Khách lẻ"`, '
           '`giam_gia = 0`. Gọi bảng kết quả là `df`.',
           'Tạo `cuoc = phi_mo_khoa + phut × gia_phut` và `doanh_thu = cuoc × (1 − giam_gia)`.'),
        p('In ra: số dòng của `df`; số chuyến “Khách lẻ”; tổng doanh thu (có dấu phân cách hàng nghìn); doanh thu '
          'theo thành phố và theo tên trạm (triệu đồng, 1 chữ số thập phân, giảm dần).'),

        h2('Phần D — Biến đổi & thu gọn dữ liệu', '1,5 điểm'),
        p('Dùng bảng `df` đã tạo ở câu C3.'),
        h3('Câu D1. Biến đổi dữ liệu', '0,75 điểm'),
        ol('Mã hoá thứ bậc `danh_gia` thành cột `diem_dg`: Rất không hài lòng 1 · Không hài lòng 2 · Bình thường 3 '
           '· Hài lòng 4 · Rất hài lòng 5; điền ô trống bằng mode, kiểu số nguyên.',
           'Tạo cột `log_phut = log(1 + phut)`; in độ lệch (skewness, `scipy.stats.skew`) của `phut` trước và sau '
           'biến đổi.',
           'Rời rạc hoá giờ mở khoá thành cột `khung_gio` bằng `pd.cut`: Sáng [0h, 9h) · Trưa chiều [9h, 16h) · '
           'Tan tầm [16h, 20h) · Tối [20h, 24h) — mút trái thuộc khung, mút phải không. In số chuyến mỗi khung; '
           'mã hoá one-hot `khung_gio` với tiền tố `kg`, in tên các cột mới.'),
        h3('Câu D2. Thu gọn dữ liệu theo khách hàng', '0,75 điểm'),
        ol('Chỉ xét khách có gói cước (bỏ “Khách lẻ”). Tạo bảng `khach`, mỗi khách một dòng, gồm: `so_chuyen` '
           '(số chuyến), `tong_chi` (tổng doanh thu, **nghìn đồng**), `phut_tb`, `km_tb`, `diem_tb` (trung bình '
           'thời lượng, quãng đường, điểm đánh giá), `ti_le_xe_may` (tỉ lệ chuyến đi bằng Xe máy điện). '
           'In kích thước bảng.',
           'Chia `tong_chi` thành 4 nhóm **cùng số lượng** Thấp / Trung bình / Khá / Cao; in số khách mỗi nhóm.',
           'Tính ma trận tương quan 6 cột số; chỉ ra các cặp có |r| > 0,7 và giải thích vì sao có thể coi một '
           'cột trong mỗi cặp là dư thừa.',
           'Chuẩn hoá z-score 6 cột số, chạy PCA; in số thành phần ít nhất cần để giữ ≥ 90% phương sai.'),

        h2('Phần E — Trực quan hoá', '2,0 điểm'),
        p('Vẽ **một hình gồm 4 biểu đồ** (2 hàng × 2 cột):'),
        ul('Trên trái: biểu đồ đường doanh thu theo tháng của **từng loại xe** (3 đường), có điểm đánh dấu và '
           'chú thích.',
           'Trên phải: biểu đồ phân tán `quang_duong_km` theo `phut`, tô màu theo loại xe.',
           'Dưới trái: biểu đồ **cột ghép** (không chồng) số chuyến theo `khung_gio`, mỗi khung giờ có một cột '
           'cho từng thành phố; các khung giờ xếp đúng thứ tự Sáng → Tối.',
           'Dưới phải: heatmap ma trận tương quan của câu D2, có ghi số trên từng ô, thang màu từ −1 đến 1.'),
        p('Mỗi biểu đồ có tiêu đề và nhãn cho các trục, ghi rõ đơn vị. Lưu hình thành `bieu_do.png` với độ phân '
          'giải 150 dpi. Viết 2 câu nhận xét: (1) xu hướng doanh thu của Xe máy điện so với hai loại xe còn lại; '
          '(2) biểu đồ phân tán cho thấy gì về tốc độ di chuyển của ba loại xe.'),
        ('end', 'Cán bộ coi thi không giải thích gì thêm.'),
    ]


def bang_a2():
    """Bảng kết quả A2 tính trực tiếp từ dữ liệu (tránh gõ nhầm 24 dòng số)."""
    mau = re.compile(r'(\d{4}-\d{2}-\d{2}) (\d{2}:\d{2}) \| ((?:DN|HA)\d{2}) \| (X[DM]\d{3}) '
                     r'\| (MUON|TRA) \| pin=(\d{1,3})%')
    so_tra, so_yeu = Counter(), Counter()
    for dong in (GOC / 'du_lieu' / 'nhat_ky_khoa_xe.txt').read_text(encoding='utf-8').splitlines():
        m = mau.fullmatch(dong)
        if m and m.group(5) == 'TRA':
            so_tra[m.group(4)] += 1
            so_yeu[m.group(4)] += int(m.group(6)) < 20
    du = sorted(x for x in so_tra if so_tra[x] >= 8)
    thieu = sorted(x for x in so_tra if so_tra[x] < 8)
    cao_nhat = max(du, key=lambda x: so_yeu[x] / so_tra[x])

    def o(x):
        o_ = [x, str(so_tra[x]), str(so_yeu[x]), f'{so_yeu[x] / so_tra[x]:.3f}']
        return [f'**{v}**' for v in o_] if x == cao_nhat else o_
    nua = (len(du) + 1) // 2
    rows = []
    for i in range(nua):
        trai = o(du[i])
        phai = o(du[nua + i]) if nua + i < len(du) else ['', '', '', '']
        rows.append(trai + phai)
    for x in thieu:
        rows.append([x, str(so_tra[x]), str(so_yeu[x]), f'{so_yeu[x] / so_tra[x]:.3f} — bị loại (< 8 lần trả)',
                     '', '', '', ''])
    return table(['Xe', 'Lần trả', 'Pin yếu', 'Tỉ lệ'] * 2, rows, small=True)


def dap_an():
    return [
        ('header', 'ĐÁP ÁN & THANG ĐIỂM', 'LƯU HÀNH NỘI BỘ — KHÔNG PHÁT CHO SINH VIÊN'),
        ('title', f'ĐÁP ÁN — {TEN_MON.upper()}',
         f'Đề số {DE_SO} · Kiểm tra thực hành 60 phút · Buổi 1–5 · Dữ liệu sinh với SEED = 2502'),
        h2('Nguyên tắc chấm chung'),
        ul('Chấm theo **kết quả và phương pháp**. Code viết khác đáp án nhưng ra đúng kết quả vẫn được đủ điểm.',
           '**Không trừ điểm dây chuyền**: nếu câu trước sai (vd. C2) làm số liệu câu sau lệch, câu sau vẫn được đủ '
           'điểm khi phương pháp đúng trên dữ liệu của chính sinh viên.',
           'Chấp nhận sai lệch làm tròn ở chữ số cuối (±0,1 triệu đồng; ±0,001 với tỉ lệ và độ lệch).',
           'Ô code báo lỗi khi Restart & Run All: không chấm phần kết quả của ô đó.',
           'Câu trả lời bằng lời (A2, D2, E) phải gắn với bối cảnh bài toán, không chỉ nêu con số.',
           'Số liệu trong các bảng dưới đây in đúng như notebook `dap_an/dap_an.ipynb` (dấu chấm thập phân).'),
        table(['Phần', 'Nội dung', 'Câu', 'Điểm'], [
            ['A', 'Python thuần: tập tin văn bản, regex, hàm', 'A1, A2', '2,0'],
            ['B', 'NumPy & lập trình hướng đối tượng', 'B1, B2, B3', '2,0'],
            ['C', 'Pandas: đọc, làm sạch, tích hợp dữ liệu', 'C1, C2, C3', '2,5'],
            ['D', 'Biến đổi & thu gọn dữ liệu', 'D1, D2', '1,5'],
            ['E', 'Trực quan hoá', 'E', '2,0'],
            ['**Tổng**', '', '', '**10,0**'],
        ], widths=[10, 60, 18, 12]),

        # ------------------------------------------------------------------ A
        h2('Phần A — Python thuần: tập tin văn bản, regex, hàm', '2,0 điểm'),
        h3('Câu A1. Lọc dòng hợp lệ bằng regex', '1,0 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['Số dòng hợp lệ', '754'],
            ['Số dòng không hợp lệ', '48 (6 kiểu lỗi × 8 dòng)'],
            ['Số sự kiện theo trạm', 'DN01 132 · DN02 135 · DN03 119 · DN04 122 · HA01 136 · HA02 110'],
            ['`dong_loi.txt`', '48 dòng'],
            ['Dòng bỏ qua', '7 dòng bắt đầu bằng # · 12 dòng trống'],
        ], widths=[35, 65]),
        code(r'''
MAU = re.compile(r'(\d{4}-\d{2}-\d{2}) (\d{2}:\d{2}) \| ((?:DN|HA)\d{2}) \| (X[DM]\d{3}) '
                 r'\| (MUON|TRA) \| pin=(\d{1,3})%')
hop_le, loi = [], []
with open(DATA_DIR / 'nhat_ky_khoa_xe.txt', encoding='utf-8') as f:
    for dong in f:
        dong = dong.rstrip('\n')
        if dong.strip() == '' or dong.startswith('#'):
            continue
        (hop_le if MAU.fullmatch(dong) else loi).append(dong)
with open('dong_loi.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(loi) + '\n')
dem_tram = Counter(MAU.fullmatch(d).group(3) for d in hop_le)
print(len(hop_le), len(loi), dict(sorted(dem_tram.items())))
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Mẫu regex đúng: thoát ký tự `|`, `(DN|HA)\\d{2}`, `X[DM]\\d{3}`, nhóm MUON/TRA, `pin=\\d{1,3}%`', '0,3'],
            ['Khớp toàn dòng (`fullmatch` hoặc `^…$`)', '0,1'],
            ['Bỏ qua dòng trống và dòng #', '0,1'],
            ['Số dòng hợp lệ / không hợp lệ đúng', '0,2'],
            ['Đếm sự kiện theo trạm bằng `Counter`, in theo thứ tự mã trạm', '0,1'],
            ['Ghi `dong_loi.txt` đúng (utf-8, mỗi dòng lỗi một dòng)', '0,2'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', 'Dùng `re.match` hoặc `re.search` với mẫu không có `$` → 8 dòng thừa trường '
             '`| gps=off` ở cuối lọt vào, ra **762** dòng hợp lệ và 40 dòng lỗi — mất 0,1 (khớp toàn dòng) và 0,2 '
             '(số dòng). Không thoát `|` → regex hiểu là phép “hoặc”: `fullmatch` ra 0 dòng hợp lệ, `search` khớp '
             'cả 802 dòng. Tính cả dòng trống và dòng # vào lỗi → ra 67 thay vì 48. Quên `rstrip(\'\\n\')` khi '
             'dùng `fullmatch` → 0 dòng hợp lệ. 6 kiểu lỗi cài sẵn: mã xe 2 chữ số (`XD08`) · sự kiện viết '
             'thường · thiếu trường pin · giờ `07h15` · trạm ngoài hệ thống `QN01` · thừa trường cuối dòng.'),
        h3('Câu A2. Thống kê pin yếu', '1,0 điểm'),
        bang_a2(),
        p('Xe có tỉ lệ pin yếu cao nhất: **XM003 — 0.350 (35,0%)**. Có sự kiện ở cả hai thành phố: '
          '**XD004, XD012, XM008, XM010**.'),
        code('''
def ti_le_pin_yeu(cac_dong, nguong=20, toi_thieu=8):
    so_tra, so_yeu = Counter(), Counter()
    for d in cac_dong:
        _, _, _, xe, su_kien, pin = MAU.fullmatch(d).groups()
        if su_kien != 'TRA':
            continue
        so_tra[xe] += 1
        if int(pin) < nguong:
            so_yeu[xe] += 1
    return {xe: so_yeu[xe] / n for xe, n in so_tra.items() if n >= toi_thieu}

kq = ti_le_pin_yeu(hop_le)
max(kq, key=kq.get)                                        # 'XM003'

thanh_pho = {}
for d in hop_le:
    _, _, tram, xe, *_ = MAU.fullmatch(d).groups()
    thanh_pho.setdefault(xe, set()).add(tram[:2])
sorted(xe for xe, s in thanh_pho.items() if s == {'DN', 'HA'})   # XD004, XD012, XM008, XM010
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Hàm đúng tên, tham số, giá trị mặc định `nguong=20`, `toi_thieu=8`', '0,15'],
            ['Chỉ xét sự kiện TRA; đổi pin sang `int` và so sánh **nhỏ hơn hẳn** `nguong`', '0,25'],
            ['Đếm theo xe, giữ xe có **≥** `toi_thieu` lần trả, trả về dict', '0,2'],
            ['In tỉ lệ 3 chữ số thập phân, xác định đúng XM003', '0,2'],
            ['Dùng `set` tìm đúng 4 xe có sự kiện ở cả hai thành phố', '0,2'],
        ], widths=[88, 12]),
        note('Bẫy:', 'Dùng `<=` thay `<`: có 14 lần trả đúng 20%, riêng XD007 có 4 lần → XD007 vọt lên **0.500** '
             'và thành xe cao nhất — mất 0,25 (so sánh) và 0,2 (kết luận). Quên lọc TRA (đếm cả MUON): XM010 có 10 '
             'sự kiện nên lọt qua ngưỡng 8, tỉ lệ 0.400 và thành xe cao nhất — mất 0,25 và 0,2. Không áp dụng '
             '`toi_thieu`: XM010 có 0.800 nhưng chỉ 5 lần trả. So sánh chuỗi `pin < \'20\'` (quên `int`) → '
             '`\'5\' < \'20\'` là False, kết quả ra XD007 0.182 — mất 0,25. Dùng `> toi_thieu` làm mất XM009 '
             '(đúng 8 lần trả) — trừ 0,1. Câu thành phố xét **tất cả** dòng hợp lệ nên phải có XM010; thiếu XM010 '
             'trừ 0,1.'),

        # ------------------------------------------------------------------ B
        h2('Phần B — NumPy & lập trình hướng đối tượng', '2,0 điểm'),
        h3('Câu B1. Ma trận lượt thuê', '0,5 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['Tổng lượt theo trạm', 'DN01 215 · DN02 211 · DN03 209 · DN04 123 · HA01 219 · HA02 170 → cao nhất '
                                    '**HA01**'],
            ['Trung bình theo khung giờ', '05–08h 31.33 · 08–11h 34.33 · 11–14h 29.67 · 14–17h 36.33 · 17–20h 59.50 '
                                          '→ cao nhất **17–20h**'],
            ['Tỉ trọng giờ cao điểm', 'DN01 0.526 · DN02 0.588 · DN03 0.368 · DN04 0.447 · HA01 0.434 · HA02 0.476 '
                                      '→ lớn hơn 0,5: **DN01, DN02**'],
            ['Số ô lớn hơn TB của chính khung giờ', '**14**'],
        ], widths=[30, 70]),
        code('''
tong_tram = L.sum(axis=1);  TRAM[tong_tram.argmax()]      # 'HA01'
tb_khung = L.mean(axis=0);  KHUNG[tb_khung.argmax()]      # '17–20h'
ty_trong = L / L.sum(axis=1, keepdims=True)                # (6, 5) / (6, 1)
cao_diem = ty_trong[:, [0, 4]].sum(axis=1)                 # chọn cột 0 và cột 4
np.array(TRAM)[cao_diem > 0.5]                             # ['DN01', 'DN02']
(L > L.mean(axis=0)).sum()                                 # 14 — (6, 5) so với (5,)
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Đúng `axis` cho tổng/trung bình và dùng `argmax`', '0,2'],
            ['Tỉ trọng bằng broadcasting, chọn đúng 2 cột không liền nhau, lọc trạm bằng mặt nạ boolean', '0,15'],
            ['Mặt nạ so với trung bình **từng cột**, ra 14', '0,15'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', 'So với trung bình của từng trạm (`axis=1`) ra 12, so với trung bình cả ma trận ra '
             '13 → mất 0,15. Viết lát cắt `ty_trong[:, 0:4]` thay vì chọn cột `[:, [0, 4]]` → cộng nhầm 4 khung đầu '
             '(`[:, [0, -1]]` hay `[:, ::4]` vẫn đúng). Chia (6, 5) '
             'cho `L.sum(axis=1)` shape (6,) → `ValueError` broadcasting. Tổng DN01 (215) sát HA01 (219): đọc nhầm '
             'dòng/cột sẽ ra DN01.'),
        h3('Câu B2. Giải hệ phương trình', '0,5 điểm'),
        p('Phí mở khoá **5 nghìn** · giá mỗi phút **0,6 nghìn** · phụ phí mỗi km **2 nghìn**. det(A) = −60 ≠ 0 nên '
          'hệ có nghiệm duy nhất. Dự đoán chuyến 30 phút, 6 km: 5 + 0,6 × 30 + 2 × 6 = **35 nghìn đồng**.'),
        code('''
A = np.array([[1, 20, 3], [1, 35, 8], [1, 50, 9]])      # cột 1 là hệ số của phí mở khoá
b = np.array([23, 42, 53])
x = np.linalg.solve(A, b)                  # [5.  0.6 2. ]
np.allclose(A @ x, b)                      # True
np.array([1, 30, 6]) @ x                   # 35.0
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Lập đúng ma trận hệ số A (3×3, có cột hệ số 1 cho phí mở khoá) và vector b', '0,15'],
            ['`np.linalg.solve` ra đúng nghiệm', '0,2'],
            ['Kiểm tra lại bằng `A @ x` (hoặc `np.dot`) và `np.allclose`', '0,1'],
            ['Dự đoán đúng 35 nghìn đồng bằng tích vô hướng với nghiệm', '0,05'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', 'Bỏ cột hệ số 1 → ma trận 3×2, `np.linalg.solve` báo `LinAlgError` (ma trận '
             'không vuông) — không có điểm hai tiêu chí đầu. Dự đoán bằng `[30, 6]` quên phí mở khoá → 30 nghìn, '
             'mất 0,05. Dùng `np.linalg.inv(A) @ b` vẫn đúng kết quả — cho đủ điểm.'),
        h3('Câu B3. Lớp chuyến đi', '1,0 điểm'),
        code('''
CD001: 25 phút · 6.0 km · 14.4 km/h · cước 17,500đ
CD002: 48 phút · 10.4 km · 13.0 km/h · cước 9,000đ
Lỗi: CD001: thời lượng 0 phút không hợp lệ
'''),
        code('''
class ChuyenDi:
    PHI_MO_KHOA = 5000

    def __init__(self, ma, phut, km, gia_phut):
        self.ma, self.km, self.gia_phut = ma, km, gia_phut
        self.phut = phut                                   # đi qua setter

    @property
    def phut(self):
        return self._phut

    @phut.setter
    def phut(self, gia_tri):
        if gia_tri <= 0:
            raise ValueError(f'{self.ma}: thời lượng {gia_tri} phút không hợp lệ')
        self._phut = gia_tri

    def cuoc_phi(self):
        return self.PHI_MO_KHOA + self.phut * self.gia_phut

    def toc_do_tb(self):
        return self.km / (self.phut / 60)

    def __str__(self):
        return (f'{self.ma}: {self.phut} phút · {self.km:.1f} km · '
                f'{self.toc_do_tb():.1f} km/h · cước {self.cuoc_phi():,.0f}đ')

class ChuyenDiGoiThang(ChuyenDi):
    def __init__(self, ma, phut, km, gia_phut, phut_mien_phi=30):
        super().__init__(ma, phut, km, gia_phut)
        self.phut_mien_phi = phut_mien_phi

    def cuoc_phi(self):
        return max(0, self.phut - self.phut_mien_phi) * self.gia_phut
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['`@property` + `@phut.setter` ném `ValueError`; kiểm tra có hiệu lực cả lúc khởi tạo', '0,35'],
            ['Thuộc tính lớp `PHI_MO_KHOA`; `cuoc_phi`, `toc_do_tb`, `__str__` đúng định dạng mẫu', '0,25'],
            ['Lớp con dùng `super().__init__`, ghi đè `cuoc_phi` có `max(0, …)` — `__str__` của lớp cha tự dùng '
             'cước mới', '0,25'],
            ['Ba lệnh kiểm thử in đúng, bắt lỗi bằng `try/except`', '0,15'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', 'Trong `__init__` gán `self._phut = phut` trực tiếp → bỏ qua kiểm tra, trừ 0,15. '
             'Getter trả `self.phut` → `RecursionError`. Gán `self.ma` **sau** `self.phut` → setter báo '
             '`AttributeError` khi tạo thông báo lỗi. Thiếu `max(0, …)`: chuyến gói tháng 20 phút ra cước âm '
             '(−5,000đ) — kiểm thử CD002 không lộ lỗi nhưng trừ 0,1. Viết cứng `5000` trong `cuoc_phi` thay vì '
             'dùng thuộc tính lớp — trừ 0,05. Lớp con viết lại `__str__` thay vì chỉ ghi đè `cuoc_phi` vẫn ra đúng '
             'chuỗi — cho 0,15/0,25 vì không tận dụng đa hình.'),

        # ------------------------------------------------------------------ C
        h2('Phần C — Pandas: đọc, làm sạch và tích hợp dữ liệu', '2,5 điểm'),
        h3('Câu C1. Đọc dữ liệu từ ba loại nguồn', '0,5 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['chuyen_di.shape', '(2635, 9)'],
            ['khach_hang.shape · goi_cuoc.shape', '(508, 5) · (3, 2)'],
            ['tram.shape · loai_xe.shape', '(6, 4) · (3, 5)'],
            ['Giá trị thiếu trong chuyen_di', 'ket_thuc 48 · ma_kh 913 · quang_duong_km 40 · danh_gia 559'],
        ], widths=[45, 55]),
        code('''
cd = pd.read_csv(DATA_DIR / 'chuyen_di.csv')
for cot in ['bat_dau', 'ket_thuc']:
    cd[cot] = pd.to_datetime(cd[cot], format='%d/%m/%Y %H:%M')
kh = pd.read_excel(DATA_DIR / 'khach_hang.xlsx', sheet_name='khach_hang')
goi = pd.read_excel(DATA_DIR / 'khach_hang.xlsx', sheet_name='goi_cuoc')
with sqlite3.connect(DATA_DIR / 'he_thong.db') as con:
    tram = pd.read_sql('SELECT * FROM tram', con)
    loai = pd.read_sql('SELECT * FROM loai_xe', con)
cd.isna().sum()[lambda s: s > 0]
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['`read_csv` + `to_datetime` cho **cả hai** cột với `format` hoặc `dayfirst=True`', '0,15'],
            ['Đọc đúng cả hai sheet (`sheet_name`)', '0,1'],
            ['`sqlite3.connect` + `pd.read_sql` với câu lệnh SELECT cho 2 bảng', '0,15'],
            ['In shape 5 bảng và số giá trị thiếu', '0,1'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', '`pd.to_datetime` không chỉ định định dạng → pandas đoán tháng/ngày từ dòng đầu '
             '(01/03/2025) rồi báo `ValueError` ở “13/03/2025 07:15”. Dùng `format=\'mixed\'` thì không báo lỗi '
             'nhưng đảo ngày–tháng với mọi ngày ≤ 12: dữ liệu rải khắp tháng 1–12 và biểu đồ Phần E sai hình. '
             'Trừ 0,15. Chỉ chuyển `bat_dau` → bước 3 của C2 báo `TypeError` khi trừ chuỗi.'),
        h3('Câu C2. Làm sạch bảng chuyến đi', '1,0 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['Sau bước 1 (bỏ trùng)', '2600 dòng — bỏ 35'],
            ['Sau bước 2', 'Từ 15 cách viết còn đúng 3 giá trị'],
            ['Bước 3', '9 giá trị `phut` ≤ 0 thành NaN (cùng 48 chuyến trống `ket_thuc` → 57 NaN) · bỏ 27 chuyến '
                       'mở khoá thử → **2573 dòng**'],
            ['Bước 4', '11 giá trị GPS lỗi thành NaN (5 âm, 6 lớn hơn 80)'],
            ['Bước 5 — ngưỡng trên', 'Xe đạp 44.5 · Xe đạp điện 57.0 · Xe máy điện 96.0 phút'],
            ['Bước 5 — ngoại lai', 'Xe đạp 31 · Xe đạp điện 37 · Xe máy điện 23'],
            ['Số chuyến sau làm sạch', 'Xe đạp 881 · Xe đạp điện 976 · Xe máy điện 716'],
            ['Trung vị phut theo loại xe', 'Xe đạp 17 · Xe đạp điện 24 · Xe máy điện 37'],
        ], widths=[30, 70]),
        code('''
cd = cd.drop_duplicates()                                                        # 1
CHUAN = {'xe đạp': 'Xe đạp', 'xe dap': 'Xe đạp',
         'xe đạp điện': 'Xe đạp điện', 'xe dap dien': 'Xe đạp điện',
         'xe máy điện': 'Xe máy điện', 'xe may dien': 'Xe máy điện'}
cd['loai_xe'] = cd['loai_xe'].str.strip().str.lower().map(CHUAN)                 # 2
cd['phut'] = (cd['ket_thuc'] - cd['bat_dau']).dt.total_seconds() / 60           # 3
cd.loc[cd['phut'] <= 0, 'phut'] = np.nan
cd = cd[~(cd['phut'] < 2)]                                                       #   giữ NaN
sai_gps = (cd['quang_duong_km'] < 0) | (cd['quang_duong_km'] > 80)             # 4
cd.loc[sai_gps, 'quang_duong_km'] = np.nan
q1 = cd.groupby('loai_xe')['phut'].transform(lambda s: s.quantile(0.25))         # 5
q3 = cd.groupby('loai_xe')['phut'].transform(lambda s: s.quantile(0.75))
tren = q3 + 1.5 * (q3 - q1)
cd.loc[cd['phut'] > tren, 'loai_xe'].value_counts()
cd['phut'] = cd['phut'].clip(upper=tren)
for cot in ['phut', 'quang_duong_km']:
    cd[cot] = cd[cot].fillna(cd.groupby('loai_xe')[cot].transform('median'))
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Bước 1: `drop_duplicates()`', '0,1'],
            ['Bước 2: `strip` + `lower` + ánh xạ **khớp nguyên chuỗi**, cả bản không dấu; còn đúng 3 giá trị', '0,2'],
            ['Bước 3: `phut` bằng `dt.total_seconds() / 60`; ≤ 0 thành NaN; bỏ < 2 nhưng giữ NaN', '0,25'],
            ['Bước 4: đổi giá trị âm và > 80 thành NaN', '0,1'],
            ['Bước 5: IQR theo nhóm (`groupby` + `transform`), `clip` về ngưỡng trên, điền trung vị theo loại xe '
             'cho cả hai cột', '0,35'],
        ], widths=[88, 12]),
        note('Bẫy:', '**Lọc bằng `cd[cd[\'phut\'] >= 2]`** làm mất luôn 57 chuyến đang NaN → còn 2516 dòng thay vì '
             '2573: bước 3 tối đa 0,1. **Dùng `.dt.seconds`** thay `total_seconds()`: khoảng thời gian âm bị quy '
             'thành 1352–1422 phút, 9 lỗi đồng hồ không bị bắt ở bước 3 mà thành “ngoại lai” ở bước 5 — trừ 0,1. '
             '**Chuẩn hoá bằng `str.contains(\'đạp\')`** gộp luôn xe đạp điện vào xe đạp (1796 dòng thay vì 893): '
             'bước 2 không có điểm; chỉ `strip().str.capitalize()` còn 6 giá trị (bản không dấu tách riêng) — '
             'cũng không có điểm. **IQR chung cả cột**: ngưỡng 66,0 phút đánh dấu 103 chuyến xe máy điện (gần 15% '
             'loại này) chỉ vì xe máy điện vốn đi lâu hơn — bước 5 tối đa 0,1.'),
        h3('Câu C3. Tích hợp dữ liệu', '1,0 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['`gia_30_phut` là cột dẫn xuất', '`(loai[\'gia_30_phut\'] == loai[\'phi_mo_khoa\'] + 30 * '
                                               'loai[\'gia_phut\']).all()` → True'],
            ['Mã khách trùng', '8 (3 dòng trùng hoàn toàn + 5 dòng chỉ khác cách viết số điện thoại) → còn 500 '
                               'khách'],
            ['Số điện thoại không đủ 10 chữ số', '9'],
            ['Khách theo gói', 'Tiêu chuẩn 261 · Sinh viên 156 · Thân thiết 83'],
            ['Chuyến có tram_muon ngoài danh mục', '24 (đều là DN07)'],
            ['Số dòng df · chuyến Khách lẻ', '2549 · 918 (gồm ô trống và mã khách không tồn tại)'],
            ['Tổng doanh thu', '55,701,300 VNĐ'],
            ['Doanh thu theo thành phố (triệu)', 'Đà Nẵng 33.3 · Hội An 22.4'],
            ['Doanh thu theo trạm (triệu)', 'Phố cổ 13.6 · Cầu Rồng 11.3 · Mỹ Khê 9.1 · An Bàng 8.8 · Chợ Hàn 8.0 · '
                                            'Sơn Trà 4.9'],
        ], widths=[33, 67]),
        code(r'''
(loai['gia_30_phut'] == loai['phi_mo_khoa'] + 30 * loai['gia_phut']).all()     # True
loai = loai.drop(columns='gia_30_phut')

kh['ma_kh'].duplicated().sum()                                                 # 8
kh = kh.drop_duplicates(subset='ma_kh')
kh['so_dien_thoai'] = (kh['so_dien_thoai'].str.replace(r'\D', '', regex=True)
                                          .str.replace(r'^84', '0', regex=True))
(kh['so_dien_thoai'].str.len() != 10).sum()                                    # 9
kh = kh.merge(goi, on='goi_cuoc', how='left')

cd = cd[cd['tram_muon'].isin(tram['id'])]                                      # bỏ 24
df = (cd.merge(loai, left_on='loai_xe', right_on='ten_loai', how='left')
        .merge(kh[['ma_kh', 'goi_cuoc', 'giam_gia']], on='ma_kh', how='left',
               validate='many_to_one')
        .merge(tram, left_on='tram_muon', right_on='id', how='left')
        .drop(columns=['ten_loai', 'id']))
df['goi_cuoc'] = df['goi_cuoc'].fillna('Khách lẻ')
df['giam_gia'] = df['giam_gia'].fillna(0)
df['cuoc'] = df['phi_mo_khoa'] + df['phut'] * df['gia_phut']
df['doanh_thu'] = df['cuoc'] * (1 - df['giam_gia'])
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Chứng minh cột dẫn xuất bằng phép so sánh trên mọi dòng rồi xoá', '0,15'],
            ['Khử trùng theo `ma_kh`, chuẩn hoá số điện thoại bằng regex, đếm đúng 9, ghép gói để có `giam_gia`',
             '0,3'],
            ['Loại 24 chuyến DN07; ghép bằng `left_on`/`right_on`, `how=\'left\'` giữ đủ 2549 dòng; gán Khách lẻ '
             'và giảm giá 0', '0,3'],
            ['`cuoc`, `doanh_thu`, tổng và doanh thu theo thành phố, theo trạm đúng', '0,25'],
        ], widths=[88, 12]),
        p('Đối chiếu nhanh tổng doanh thu sai — mỗi con số chỉ ra đúng một lỗi ở phía trước:'),
        table(['Sinh viên ra', 'Nguyên nhân'], [
            ['56,460,570 (2583 dòng)', 'Không bỏ dòng trùng (C2 bước 1)'],
            ['54,521,960 (2492 dòng)', 'Lọc `phut >= 2` làm mất các chuyến đang trống `phut` (C2 bước 3)'],
            ['56,551,040', 'Không winsorize ngoại lai (C2 bước 5)'],
            ['54,371,880', 'Tính IQR chung cả cột thay vì theo loại xe (C2 bước 5)'],
            ['55,505,540', 'Điền trung vị chung thay vì theo loại xe (C2 bước 5)'],
            ['55,955,650 (2565 dòng)', 'Khử trùng khách bằng `drop_duplicates()` không có `subset` → chỉ bỏ được 3, '
                                       '5 khách còn trùng nhân đôi 16 chuyến (C3). `validate=\'many_to_one\'` sẽ báo '
                                       'lỗi ngay'],
            ['33,737,600 (1631 dòng)', '`how=\'inner\'` khi ghép khách hàng — mất toàn bộ khách lẻ'],
            ['58,795,200', 'Quên trừ giảm giá'],
            ['56,237,920 (2573 dòng)', 'Không loại trạm DN07 — tổng tiền vẫn tính được nhưng 24 chuyến không có '
                                       '`ten_tram`, `thanh_pho`'],
        ], widths=[28, 72]),
        note('Lỗi thường gặp khác:', 'Chỉ bỏ khoảng trắng và thay `+84` → còn 68 số “không đủ 10 chữ số” (sót '
             'dấu chấm, gạch ngang, đầu `84-`); bỏ ký tự không phải số nhưng quên đổi `84` → 71. Ép '
             '`so_dien_thoai` sang số nguyên làm mất số 0 ở đầu. Ghép trạm bằng `on=\'tram_muon\'` → `KeyError` vì '
             'bảng `tram` không có cột này.'),

        # ------------------------------------------------------------------ D
        h2('Phần D — Biến đổi & thu gọn dữ liệu', '1,5 điểm'),
        h3('Câu D1. Biến đổi dữ liệu', '0,75 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['Mode của diem_dg', '4 (Hài lòng)'],
            ['Phân bố diem_dg sau khi điền', '1: 2 · 2: 95 · 3: 520 · 4: 1464 · 5: 468'],
            ['Skew phut → log_phut', '1.491 → 0.020'],
            ['Số chuyến theo khung giờ', 'Sáng 718 · Trưa chiều 759 · Tan tầm 867 · Tối 205'],
            ['Cột one-hot', 'kg_Sáng · kg_Trưa chiều · kg_Tan tầm · kg_Tối'],
        ], widths=[35, 65]),
        code('''
THU_TU = {'Rất không hài lòng': 1, 'Không hài lòng': 2, 'Bình thường': 3,
          'Hài lòng': 4, 'Rất hài lòng': 5}
df['diem_dg'] = df['danh_gia'].map(THU_TU)
df['diem_dg'] = df['diem_dg'].fillna(df['diem_dg'].mode()[0]).astype(int)

df['log_phut'] = np.log1p(df['phut'])
st.skew(df['phut']), st.skew(df['log_phut'])                    # 1.491, 0.020

df['khung_gio'] = pd.cut(df['bat_dau'].dt.hour, bins=[0, 9, 16, 20, 24], right=False,
                         labels=['Sáng', 'Trưa chiều', 'Tan tầm', 'Tối'])
pd.get_dummies(df['khung_gio'], prefix='kg', dtype=int).columns.tolist()
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Ánh xạ đúng thứ tự 1–5, điền mode (`mode()[0]`), kiểu số nguyên', '0,25'],
            ['`np.log1p` (hoặc `np.log(1 + x)`), in skew trước và sau', '0,2'],
            ['`pd.cut` đúng mốc và `right=False`; one-hot đúng 4 cột có tiền tố', '0,3'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', 'Quên `right=False` → các chuyến bắt đầu lúc 9h, 16h, 20h bị xếp lùi một khung: '
             'Sáng 825 · Trưa chiều 844 · Tan tầm 779 · Tối 101 — trừ 0,15. Dùng `astype(\'category\').cat.codes` '
             '→ mã theo thứ tự chữ cái (Bình thường 0, Hài lòng 1, Không hài lòng 2, …) sai thứ bậc: trừ 0,2. '
             'Quên `[0]` sau `mode()` → phần lớn ô vẫn trống. Dùng `np.log` không cộng 1 vẫn chạy (skew −0.054) vì '
             'sau C2 mọi `phut` ≥ 2 — chấp nhận nhưng nhắc rủi ro.'),
        h3('Câu D2. Thu gọn dữ liệu theo khách hàng', '0,75 điểm'),
        table(['Chỉ số', 'Kết quả chuẩn'], [
            ['Kích thước bảng khach', '(474, 6) — trong 500 khách có 26 khách không có chuyến nào'],
            ['Nhóm chi tiêu (qcut)', 'Thấp 120 · Trung bình 117 · Khá 118 · Cao 119 (474 không chia hết cho 4 và có '
                                     'giá trị trùng ở mép nhóm)'],
            ['Cặp |r| > 0,7', 'phut_tb – km_tb **0.91** · so_chuyen – tong_chi **0.78**'],
            ['Tỉ lệ phương sai tích luỹ', '0.508 · 0.725 · 0.885 · **0.973** · 0.988 · 1.000'],
            ['Số thành phần PCA giữ ≥ 90%', '**4**'],
        ], widths=[30, 70]),
        code('''
thanh_vien = df[df['goi_cuoc'] != 'Khách lẻ']
khach = thanh_vien.groupby('ma_kh').agg(
    so_chuyen=('ma_cd', 'count'), tong_chi=('doanh_thu', 'sum'),
    phut_tb=('phut', 'mean'), km_tb=('quang_duong_km', 'mean'), diem_tb=('diem_dg', 'mean'),
    ti_le_xe_may=('loai_xe', lambda s: (s == 'Xe máy điện').mean()))
khach['tong_chi'] = khach['tong_chi'] / 1e3
khach['nhom_chi'] = pd.qcut(khach['tong_chi'], q=4,
                            labels=['Thấp', 'Trung bình', 'Khá', 'Cao'])

tq = khach.drop(columns='nhom_chi').corr()
Z = StandardScaler().fit_transform(khach.drop(columns='nhom_chi'))
PCA(n_components=0.9).fit(Z).n_components_                      # 4
'''),
        p('Giải thích được chấp nhận: `km_tb` gần như bằng tốc độ × `phut_tb` ÷ 60 — quãng đường tỉ lệ thuận với '
          'thời gian đi nên hai cột mang cùng thông tin về “độ dài chuyến”; `tong_chi` ≈ `so_chuyen` × cước trung '
          'bình mỗi chuyến nên phần lớn thông tin của `tong_chi` đã nằm trong số chuyến. Giữ một cột trong mỗi cặp '
          'là đủ cho phần lớn phân tích. PCA “gộp” được hai cặp này nên 6 cột còn 4 thành phần.'),
        table(['Tiêu chí', 'Điểm'], [
            ['`groupby` + `agg` ra đúng 6 cột, bỏ Khách lẻ, đơn vị nghìn đồng', '0,25'],
            ['`pd.qcut` 4 nhóm', '0,1'],
            ['Ma trận tương quan, chỉ ra đúng 2 cặp và giải thích dư thừa', '0,2'],
            ['Chuẩn hoá trước PCA, ra 4 thành phần (sklearn hoặc tự tính trị riêng bằng numpy)', '0,2'],
        ], widths=[88, 12]),
        note('Lỗi thường gặp:', 'Không chuẩn hoá trước PCA → `tong_chi` (phương sai 4806) lấn át các cột khác '
             '(`ti_le_xe_may` chỉ 0,09), PC1 đã giữ 98,0% → ra 1 thành phần: phần PCA 0,05. Đọc nhầm 3 thành phần '
             '(88,5%) là “đủ 90%” → trừ 0,1. Dùng `pd.cut` thay `qcut` → bốn nhóm cùng độ rộng nhưng lệch số '
             'lượng (385 · 67 · 17 · 5): không có điểm ý qcut. Để `tong_chi` ở đơn vị đồng không ảnh hưởng tương '
             'quan và PCA đã chuẩn hoá — chỉ trừ 0,05 ở ý đơn vị.'),

        # ------------------------------------------------------------------ E
        h2('Phần E — Trực quan hoá', '2,0 điểm'),
        p('Hình mẫu: `dap_an/bieu_do.png`.'),
        ('img', GOC / 'dap_an' / 'bieu_do.png', 'Hình kết quả mẫu Phần E — lưu với dpi = 150'),
        table(['Tháng', '3', '4', '5', '6', '7', '8'], [
            ['Xe đạp (triệu)', '0.9', '1.1', '1.0', '0.8', '0.9', '0.8'],
            ['Xe đạp điện (triệu)', '2.2', '2.5', '2.6', '3.0', '3.4', '2.9'],
            ['Xe máy điện (triệu)', '2.1', '2.4', '4.2', '7.3', '8.9', '8.8'],
        ], widths=[28, 12, 12, 12, 12, 12, 12], small=True),
        table(['Số chuyến', 'Sáng', 'Trưa chiều', 'Tan tầm', 'Tối'], [
            ['Đà Nẵng', '473', '478', '562', '126'],
            ['Hội An', '245', '281', '305', '79'],
        ], widths=[28, 18, 18, 18, 18], small=True),
        code('''
fig, axes = plt.subplots(2, 2, figsize=(14, 10))
THU_TU_XE = ['Xe đạp', 'Xe đạp điện', 'Xe máy điện']

theo_thang = (df.groupby([df['bat_dau'].dt.month, 'loai_xe'])['doanh_thu'].sum() / 1e6
              ).unstack()[THU_TU_XE]
theo_thang.plot(ax=axes[0, 0], marker='o')
axes[0, 0].set(title='Doanh thu theo tháng và loại xe', xlabel='Tháng (2025)',
               ylabel='Doanh thu (triệu đồng)')

sns.scatterplot(data=df, x='phut', y='quang_duong_km', hue='loai_xe', hue_order=THU_TU_XE,
                alpha=0.4, s=12, ax=axes[0, 1])
axes[0, 1].set(title='Quãng đường theo thời lượng chuyến', xlabel='Thời lượng (phút)',
               ylabel='Quãng đường (km)')

kg = pd.crosstab(df['khung_gio'], df['thanh_pho'])      # category → giữ thứ tự Sáng → Tối
kg.plot(kind='bar', ax=axes[1, 0], rot=0)
axes[1, 0].set(title='Số chuyến theo khung giờ và thành phố', xlabel='Khung giờ bắt đầu',
               ylabel='Số chuyến')

sns.heatmap(tq, annot=True, fmt='.2f', cmap='coolwarm', vmin=-1, vmax=1, ax=axes[1, 1])
axes[1, 1].set_title('Tương quan các đặc trưng khách hàng')

fig.tight_layout()
fig.savefig('bieu_do.png', dpi=150)      # lưu TRƯỚC plt.show()
plt.show()
'''),
        table(['Tiêu chí', 'Điểm'], [
            ['Một hình, 2 hàng × 2 cột (`plt.subplots(2, 2)`), mỗi biểu đồ vẽ đúng ô bằng `ax=`', '0,2'],
            ['Biểu đồ đường: 3 loại xe, có marker và chú thích, trục tháng 3–8', '0,35'],
            ['Biểu đồ phân tán phut × quang_duong_km tô màu theo loại xe', '0,3'],
            ['Cột ghép khung giờ × thành phố, **không chồng**, thứ tự Sáng → Tối', '0,35'],
            ['Heatmap có số trên ô, thang màu khoá −1 đến 1', '0,3'],
            ['Tiêu đề và nhãn trục có đơn vị — cả bốn biểu đồ', '0,2'],
            ['Lưu `bieu_do.png`, `dpi=150`', '0,1'],
            ['Hai câu nhận xét hợp lý', '0,2'],
        ], widths=[88, 12]),
        p('Nhận xét được chấp nhận: (1) doanh thu Xe máy điện tăng mạnh từ 2,1 triệu (tháng 3) lên 8,9 triệu '
          '(tháng 7), vượt Xe đạp điện từ tháng 5 và gấp khoảng 3 lần vào tháng 8, trong khi Xe đạp điện chỉ nhích '
          'nhẹ và Xe đạp đi ngang quanh 1 triệu/tháng (tỉ trọng chuyến xe máy điện tăng từ 13% lên 39%); (2) các '
          'điểm xếp thành những dải thẳng đi qua gốc toạ độ — quãng đường tỉ lệ thuận với thời lượng — với độ dốc '
          'khác nhau: xe máy điện nhanh nhất (≈ 23 km/h), xe đạp điện ≈ 16 km/h, xe đạp ≈ 11 km/h; điều này giải '
          'thích tương quan 0,91 giữa `phut_tb` và `km_tb` ở D2. Chấp nhận thêm: các cột điểm dựng đứng ở 44,5 · 57 '
          '· 96 phút là dấu vết winsorization của C2; Tan tầm là khung đông nhất ở cả hai thành phố, Đà Nẵng gần gấp '
          'đôi Hội An ở mọi khung giờ.'),
        note('Lỗi thường gặp:', 'Gọi `savefig` sau `plt.show()` → ảnh trắng, mất 0,1. Vẽ bằng `plt.plot` thay vì '
             '`ax=` khiến các biểu đồ đè lên ô cuối cùng: phần bố cục 0. `khung_gio` bị đổi sang chuỗi (hoặc tự '
             'tạo bằng if/else) → crosstab xếp theo chữ cái Sáng, Tan tầm, Trưa chiều, Tối — trừ 0,1. Dùng '
             '`stacked=True` → trừ 0,15. Heatmap không đặt `vmin=-1, vmax=1` → thang màu co về 0,02…1, mọi ô đều '
             'đỏ đậm dù tương quan yếu: trừ 0,1.'),
        ('end', 'Đáp án gồm lời giải tham khảo; mọi cách làm đúng khác đều được chấp nhận.'),
    ]


# =====================================================================================
# XUẤT HTML → PDF
# =====================================================================================
TOKEN = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*|~[^~]+~|<br>)')


def inline_html(text):
    out = []
    for phan in TOKEN.split(text):
        if not phan:
            continue
        if phan == '<br>':
            out.append('<br>')
        elif phan.startswith('`'):
            out.append(f'<code>{html.escape(phan[1:-1])}</code>')
        elif phan.startswith('**'):
            out.append(f'<b>{inline_html(phan[2:-2])}</b>')
        elif phan.startswith('~'):
            out.append(f'<i>{html.escape(phan[1:-1])}</i>')
        else:
            out.append(html.escape(phan))
    return ''.join(out)


CSS = """
@page { size: A4; margin: 18mm 17mm 18mm 17mm;
  @bottom-left { content: "%(chan)s"; font: 7pt 'Times New Roman', serif; color: #555; }
  @bottom-right { content: counter(page) " / " counter(pages); font: 7pt 'Times New Roman', serif; color: #555; } }
* { box-sizing: border-box; }
body { font-family: 'Times New Roman', Times, serif; font-size: 11.5pt; line-height: 1.55; color: #111; margin: 0; }
code, pre { font-family: Menlo, 'DejaVu Sans Mono', Consolas, monospace; }
code { font-size: 9pt; background: #f0f0f0; padding: 0 3px; border-radius: 2px; }
.hdr { display: flex; justify-content: space-between; font-weight: bold; margin-bottom: 6px; }
.hdr .l { text-align: center; } .hdr .r { text-align: center; }
.hdr .l .rule { width: 45%%; margin: 2px auto 0; border-top: 1px solid #111; }
.hdr .r .sub { font-weight: normal; }
.stamp { border: 1px solid #a33; color: #a33; padding: 2px 10px; margin-top: 4px; font-size: 10pt; letter-spacing: .3px; }
h1 { text-align: center; font-size: 17pt; margin: 14px 0 0; }
.subt { text-align: center; font-style: italic; margin-bottom: 10px; }
.info { display: grid; grid-template-columns: 1fr 1fr 1fr; border: 1px solid #555; margin: 8px 0 12px; }
.info div { padding: 5px 8px; border-right: 1px solid #bbb; border-bottom: 1px solid #bbb; font-size: 10.5pt; }
.info div:nth-child(3n) { border-right: none; } .info div:nth-last-child(-n+3) { border-bottom: none; }
.info .k { display: block; font-family: Helvetica, Arial, sans-serif; font-size: 7.5pt; color: #555; letter-spacing: .4px; }
.sv { display: flex; gap: 10px; margin: 6px 0 16px; }
.sv span { white-space: nowrap; } .sv .line { flex: 1; border-bottom: 1px dashed #333; }
h2 { display: flex; justify-content: space-between; align-items: baseline; background: #f2f2f2;
     border-left: 5px solid #1f4e79; padding: 5px 10px; font-size: 13.5pt; margin: 18px 0 8px; break-after: avoid; }
h2 .d, h3 .d { color: #1f4e79; }
h3 { display: flex; justify-content: space-between; font-size: 11.5pt; margin: 12px 0 4px; break-after: avoid; }
p { margin: 4px 0 6px; }
ul, ol { margin: 4px 0 8px; padding-left: 26px; } li { margin: 1px 0; }
pre { background: #f5f5f5; border: 1px solid #ddd; padding: 7px 10px; font-size: 8.6pt; line-height: 1.38;
      white-space: pre-wrap; margin: 6px 0 8px; }
table { width: 100%%; border-collapse: collapse; margin: 6px 0 10px; font-size: 10.5pt; break-inside: auto; }
table.small { font-size: 9.8pt; }
th, td { border: 1px solid #999; padding: 3px 7px; text-align: left; vertical-align: top; line-height: 1.35; }
th { background: #eee; } tr { break-inside: avoid; }
td code, th code { font-size: 8.4pt; }
.note { border: 1px solid #a33; background: #fbf1ef; padding: 7px 10px; margin: 6px 0 12px; font-size: 10.2pt;
        line-height: 1.45; break-inside: avoid; }
.note .n { color: #a33; font-weight: bold; }
figure { margin: 8px 0; text-align: center; border: 1px solid #ddd; padding: 6px; break-inside: avoid; }
figure img { width: 92%%; } figcaption { font-style: italic; font-size: 10pt; color: #444; }
.end { text-align: center; margin-top: 16px; } .end b { font-size: 12pt; } .end i { display: block; font-size: 10pt; }
"""


def to_html(khoi, chan, tieu_de):
    out = []
    for b in khoi:
        loai = b[0]
        if loai == 'header':
            dau = ('<div class="r"><div>{}</div><div class="sub">Học kỳ …… — Năm học 20……–20……</div>{}</div>'
                   .format(b[1], f'<div class="stamp">{b[2]}</div>' if b[2] else ''))
            out.append('<div class="hdr"><div class="l"><div>TRƯỜNG ……………………………………</div>'
                       '<div>KHOA ……………………………………</div><div class="rule"></div></div>' + dau + '</div>')
        elif loai == 'title':
            out.append(f'<h1>{html.escape(b[1])}</h1><div class="subt">{html.escape(b[2])}</div>')
        elif loai == 'info':
            out.append('<div class="info">' + ''.join(
                f'<div><span class="k">{k}</span>{html.escape(v)}</div>' for k, v in b[1]) + '</div>')
        elif loai == 'sv':
            out.append('<div class="sv"><span>Họ và tên:</span><span class="line" style="flex:3"></span>'
                       '<span>MSSV:</span><span class="line"></span><span>Lớp:</span><span class="line"></span></div>')
        elif loai in ('h2', 'h3'):
            out.append(f'<{loai}><span>{inline_html(b[1])}</span><span class="d">{b[2]}</span></{loai}>')
        elif loai == 'p':
            out.append(f'<p>{inline_html(b[1])}</p>')
        elif loai in ('ul', 'ol'):
            out.append(f'<{loai}>' + ''.join(f'<li>{inline_html(i)}</li>' for i in b[1]) + f'</{loai}>')
        elif loai == 'code':
            out.append(f'<pre>{html.escape(b[1])}</pre>')
        elif loai == 'table':
            _, header, rows, widths, small = b
            cg = ''.join(f'<col style="width:{w}%">' for w in widths) if widths else ''
            th = ''.join(f'<th>{inline_html(h)}</th>' for h in header)
            tr = ''.join('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in r) + '</tr>' for r in rows)
            out.append(f'<table class="{"small" if small else ""}"><colgroup>{cg}</colgroup>'
                       f'<thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table>')
        elif loai == 'note':
            out.append(f'<div class="note"><span class="n">{b[1]}</span> {inline_html(b[2])}</div>')
        elif loai == 'img':
            out.append(f'<figure><img src="{b[1].as_uri()}"><figcaption>{html.escape(b[2])}</figcaption></figure>')
        elif loai == 'end':
            out.append(f'<div class="end"><b>— HẾT —</b><i>{html.escape(b[1])}</i></div>')
    return (f'<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>{html.escape(tieu_de)}</title>'
            f'<style>{CSS % {"chan": chan}}</style></head><body>{"".join(out)}</body></html>')


def xuat_pdf(khoi, chan, tieu_de, dich):
    with tempfile.TemporaryDirectory() as tmp:
        f = Path(tmp) / 'tai_lieu.html'
        f.write_text(to_html(khoi, chan, tieu_de), encoding='utf-8')
        subprocess.run([CHROME, '--headless', '--disable-gpu', '--no-pdf-header-footer',
                        '--allow-file-access-from-files', f'--print-to-pdf={dich}', f.as_uri()],
                       check=True, capture_output=True)


# =====================================================================================
# XUẤT DOCX
# =====================================================================================
XANH = RGBColor(0x1F, 0x4E, 0x79)
DO = RGBColor(0xAA, 0x33, 0x33)
MONO = 'Consolas'


def to_mau(el, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    el.append(shd)


def vien_trai(par, mau='1F4E79', do_day=36):
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    for k, v in {'w:val': 'single', 'w:sz': str(do_day), 'w:space': '6', 'w:color': mau}.items():
        left.set(qn(k), v)
    bdr.append(left)
    ppr.append(bdr)


def font(run, ten='Times New Roman', co=None, dam=None, nghieng=None, mau=None):
    run.font.name = ten
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), ten)
    if co:
        run.font.size = Pt(co)
    if dam is not None:
        run.bold = dam
    if nghieng is not None:
        run.italic = nghieng
    if mau is not None:
        run.font.color.rgb = mau


def inline_docx(par, text, co=None, dam=False):
    for phan in TOKEN.split(text):
        if not phan:
            continue
        if phan == '<br>':
            par.add_run().add_break()
        elif phan.startswith('`'):
            r = par.add_run(phan[1:-1])
            font(r, MONO, (co or 12) - 2.5, dam)
        elif phan.startswith('**'):
            inline_docx(par, phan[2:-2], co, True)
        elif phan.startswith('~'):
            font(par.add_run(phan[1:-1]), co=co, dam=dam, nghieng=True)
        else:
            font(par.add_run(phan), co=co, dam=dam)


def khoang(par, truoc=0, sau=4, dong=1.15):
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(truoc), Pt(sau), dong


def tab_phai(par, doc):
    rong = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    par.paragraph_format.tab_stops.add_tab_stop(rong - Cm(0.3), WD_TAB_ALIGNMENT.RIGHT)


def bo_vien_bang(tbl):
    tblpr = tbl._tbl.tblPr
    bd = OxmlElement('w:tblBorders')
    for canh in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{canh}')
        e.set(qn('w:val'), 'nil')
        bd.append(e)
    tblpr.append(bd)


def to_docx(khoi, chan, tieu_de, dich):
    doc = Document()
    doc.core_properties.title = tieu_de
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st_ = doc.styles['Normal']
    st_.font.name = 'Times New Roman'
    st_.font.size = Pt(12)
    st_.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    fp = sec.footer.paragraphs[0]
    tab_phai(fp, doc)
    font(fp.add_run(chan + '\t'), co=8, mau=RGBColor(0x55, 0x55, 0x55))
    for truong in ['PAGE', None, 'NUMPAGES']:
        if truong is None:
            font(fp.add_run(' / '), co=8)
            continue
        r = fp.add_run()
        font(r, co=8)
        for loai_, noi_dung in [('begin', None), (None, truong), ('end', None)]:
            if loai_:
                e = OxmlElement('w:fldChar')
                e.set(qn('w:fldCharType'), loai_)
            else:
                e = OxmlElement('w:instrText')
                e.set(qn('xml:space'), 'preserve')
                e.text = noi_dung
            r._r.append(e)

    for b in khoi:
        loai = b[0]
        if loai == 'header':
            t = doc.add_table(rows=1, cols=2)
            bo_vien_bang(t)
            trai, phai = t.rows[0].cells
            for o, dong in [(trai, ['TRƯỜNG ……………………………………', 'KHOA ……………………………………']),
                            (phai, [b[1], 'Học kỳ …… — Năm học 20……–20……'])]:
                o.paragraphs[0].text = ''
                for i, d in enumerate(dong):
                    par = o.paragraphs[0] if i == 0 else o.add_paragraph()
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    khoang(par, sau=0)
                    font(par.add_run(d), co=12, dam=not (o is phai and i == 1))
            if b[2]:
                par = phai.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                font(par.add_run(b[2]), co=10, dam=True, mau=DO)
        elif loai == 'title':
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            khoang(par, truoc=10, sau=0)
            font(par.add_run(b[1]), co=16, dam=True)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            khoang(par, sau=8)
            font(par.add_run(b[2]), co=11.5, nghieng=True)
        elif loai == 'info':
            t = doc.add_table(rows=2, cols=3)
            t.style = 'Table Grid'
            for i, (k, v) in enumerate(b[1]):
                o = t.rows[i // 3].cells[i % 3]
                par = o.paragraphs[0]
                khoang(par, sau=0)
                font(par.add_run(k), 'Arial', 7, mau=RGBColor(0x55, 0x55, 0x55))
                par = o.add_paragraph()
                khoang(par, sau=2)
                font(par.add_run(v), co=10.5)
        elif loai == 'sv':
            par = doc.add_paragraph()
            khoang(par, truoc=10, sau=10)
            font(par.add_run('Họ và tên: ………………………………………………   MSSV: ………………………   Lớp: ……………………'), co=12)
        elif loai in ('h2', 'h3'):
            par = doc.add_paragraph()
            tab_phai(par, doc)
            if loai == 'h2':
                khoang(par, truoc=12, sau=6)
                to_mau(par._p.get_or_add_pPr(), 'F2F2F2')
                vien_trai(par)
                co = 13.5
            else:
                khoang(par, truoc=8, sau=3)
                co = 12
            par.paragraph_format.keep_with_next = True
            inline_docx(par, b[1], co, True)
            if b[2]:
                font(par.add_run('\t' + b[2]), co=co, dam=True, mau=XANH)
        elif loai == 'p':
            par = doc.add_paragraph()
            khoang(par)
            inline_docx(par, b[1], 12)
        elif loai in ('ul', 'ol'):
            for i, item in enumerate(b[1], 1):
                par = doc.add_paragraph()
                khoang(par, sau=2)
                par.paragraph_format.left_indent = Cm(0.9)
                par.paragraph_format.first_line_indent = Cm(-0.5)
                font(par.add_run(f'{i}.\t' if loai == 'ol' else '•\t'), co=12)
                par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
                inline_docx(par, item, 12)
        elif loai == 'code':
            par = doc.add_paragraph()
            khoang(par, truoc=2, sau=6, dong=1.0)
            to_mau(par._p.get_or_add_pPr(), 'F5F5F5')
            for i, dong in enumerate(b[1].split('\n')):
                if i:
                    par.add_run().add_break(WD_BREAK.LINE)
                font(par.add_run(dong), MONO, 8.5)
        elif loai == 'table':
            _, header, rows, widths, small = b
            co = 9.5 if small else 10.5
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            rong = 17.4
            for r_i, dong in enumerate([header] + rows):
                for c_i, gia_tri in enumerate(dong):
                    o = t.rows[r_i].cells[c_i]
                    if widths:
                        o.width = Cm(rong * widths[c_i] / sum(widths))
                    par = o.paragraphs[0]
                    khoang(par, sau=0, dong=1.0)
                    for j, phan in enumerate(gia_tri.split('<br>')):
                        if j:
                            par = o.add_paragraph()
                            khoang(par, sau=0, dong=1.0)
                        inline_docx(par, phan, co, r_i == 0)
                    if r_i == 0:
                        to_mau(o._tc.get_or_add_tcPr(), 'EEEEEE')
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif loai == 'note':
            t = doc.add_table(rows=1, cols=1)
            t.style = 'Table Grid'
            o = t.rows[0].cells[0]
            to_mau(o._tc.get_or_add_tcPr(), 'FBF1EF')
            par = o.paragraphs[0]
            khoang(par, truoc=2, sau=2, dong=1.1)
            font(par.add_run(b[1] + ' '), co=10.5, dam=True, mau=DO)
            inline_docx(par, b[2], 10.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif loai == 'img':
            doc.add_picture(str(b[1]), width=Cm(17))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(par.add_run(b[2]), co=10, nghieng=True)
        elif loai == 'end':
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            khoang(par, truoc=12, sau=0)
            font(par.add_run('— HẾT —'), co=12, dam=True)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(par.add_run(b[1]), co=10.5, nghieng=True)
    doc.save(dich)


if __name__ == '__main__':
    for khoi, chan, tieu_de, thu_muc, ten in [
        (de_thi(), f'{TEN_MON} · Đề số {DE_SO}', f'{TEN_MON} · Đề số {DE_SO}', 'de_thi', 'De_kiem_tra_LTPTDL'),
        (dap_an(), f'Đáp án & thang điểm · Đề số {DE_SO} · Lưu hành nội bộ', f'Đáp án · Đề số {DE_SO}',
         'dap_an', 'Dap_an_thang_diem'),
    ]:
        dich = GOC / thu_muc
        xuat_pdf(khoi, chan, tieu_de, dich / f'{ten}.pdf')
        to_docx(khoi, chan, tieu_de, dich / f'{ten}.docx')
        print('Đã tạo:', dich / f'{ten}.pdf', '·', f'{ten}.docx')
